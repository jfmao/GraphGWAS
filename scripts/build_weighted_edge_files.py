"""Materialize per-species weighted-edge TSVs for the M5★ weighted substrate.

For each species with a canonical-PPI list in
`src/python/graphgwas/canonical_ppi.py`, emit a TSV in the standard
weighted-edge format (gene_a, gene_b, weight, source). For human, also
fold in:
  - Billmann HAP1 qGI (data/billmann2026/billmann_hc_gi_edges.tsv)
  - DepMap ED scores (data/billmann2026/billmann_depmap_ed_edges.tsv)

Output:
  data/weighted_edges/yeast.tsv
  data/weighted_edges/arabidopsis.tsv
  data/weighted_edges/rice.tsv
  data/weighted_edges/human.tsv     ← Billmann + DepMap (NEW)

Once these exist, the M5★ kernel can consume them via
graphgwas.weighted_substrate.build_weighted_p_gg() instead of building
the binary G matrix from cache['ppi'] directly.
"""
from __future__ import annotations

import csv
import sys
import urllib.request
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src" / "python"))

# ---------------------------------------------------------------------------
# Auto-download Billmann supplementary files + extract qGI / DepMap-ED TSVs
# ---------------------------------------------------------------------------

BOONE_BASE = "https://boonelab.ccbr.utoronto.ca/supplement/billmanncostanzo2026/assets/data_files"
BILLMANN_DIR = REPO / "data" / "billmann2026"
NEEDED_XLSX = {
    "File_S4.xlsx": "Pairwise GI Complete dataset (qGI + FDR; 160 MB)",
    "File_S21.xlsx": "DepMap ED vs qGI cross-validated pairs",
}


def ensure_billmann_xlsx() -> None:
    BILLMANN_DIR.mkdir(parents=True, exist_ok=True)
    for fname, desc in NEEDED_XLSX.items():
        path = BILLMANN_DIR / fname
        if path.exists() and path.stat().st_size > 1000:
            continue
        url = f"{BOONE_BASE}/{fname}"
        print(f"Downloading {fname} ({desc}) from {url} (may take a minute)...")
        urllib.request.urlretrieve(url, path)
        print(f"  → {path.stat().st_size/1e6:.2f} MB")


def extract_qgi_tsv() -> None:
    """File_S4 'pairwise GI_Complete dataset' → billmann_hc_gi_edges.tsv."""
    out = BILLMANN_DIR / "billmann_hc_gi_edges.tsv"
    if out.exists() and out.stat().st_size > 1000:
        print(f"  qGI TSV already present: {out}")
        return
    import openpyxl
    wb = openpyxl.load_workbook(BILLMANN_DIR / "File_S4.xlsx",
                                  read_only=True, data_only=True)
    ws = wb["pairwise GI_Complete dataset"]
    print(f"  Extracting qGI: {ws.max_row:,} rows × {ws.max_column} cols ...")
    with open(out, "w") as f:
        w = csv.writer(f, delimiter="\t")
        rows_iter = ws.iter_rows(values_only=True)
        next(rows_iter)
        w.writerow(["library_gene", "query_gene", "query_screen",
                    "qGI_score", "FDR", "GI_standard", "GI_stringent"])
        for r in rows_iter:
            if r[0] is None:
                continue
            w.writerow([r[0], r[1], r[2], r[3], r[4], r[5] or "", r[6] or ""])
    wb.close()
    print(f"  Wrote {out} ({out.stat().st_size/1e6:.2f} MB)")


def extract_depmap_ed_tsv() -> None:
    """File_S21 4 sheets → combined billmann_depmap_ed_edges.tsv."""
    out = BILLMANN_DIR / "billmann_depmap_ed_edges.tsv"
    if out.exists() and out.stat().st_size > 1000:
        print(f"  DepMap ED TSV already present: {out}")
        return
    import openpyxl
    wb = openpyxl.load_workbook(BILLMANN_DIR / "File_S21.xlsx",
                                  read_only=True, data_only=True)
    sheets_to_export = [
        ("Positive ED_Negative qGI", "ED_pos_qGI_neg"),
        ("Negative ED_Positive qGI", "ED_neg_qGI_pos"),
        ("Positive ED_Positive qGI", "ED_pos_qGI_pos"),
        ("Negative ED_Negative qGI", "ED_neg_qGI_neg"),
    ]
    with open(out, "w") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["library_gene", "query_gene", "qGI_score", "qGI_FDR",
                    "GI_type", "ED_score_lib_v_query", "ED_pval_lib_v_query",
                    "ED_score_query_v_lib", "ED_pval_query_v_lib",
                    "evidence_class"])
        for sheet, evidence_class in sheets_to_export:
            for r in list(wb[sheet].iter_rows(values_only=True))[1:]:
                if r[0] is None or r[1] is None:
                    continue
                w.writerow([r[0], r[1], r[2], r[3], r[4],
                            r[5], r[6], r[7], r[8], evidence_class])
    wb.close()
    print(f"  Wrote {out} ({out.stat().st_size/1e6:.2f} MB)")


print("=== Step 0: ensure Billmann supplement files (auto-download if missing) ===")
ensure_billmann_xlsx()
extract_qgi_tsv()
extract_depmap_ed_tsv()
print()

from graphgwas.canonical_ppi import (  # noqa: E402
    YEAST_CANONICAL_EDGES,
    ARABIDOPSIS_CANONICAL_EDGES,
    RICE_CANONICAL_EDGES,
)
from graphgwas.weighted_substrate import (  # noqa: E402
    edges_from_canonical_ppi,
    edges_from_billmann_qgi_tsv,
    edges_from_depmap_ed_tsv,
    write_edges_tsv,
)

OUT = REPO / "data" / "weighted_edges"
OUT.mkdir(parents=True, exist_ok=True)
BILLMANN_DIR = REPO / "data" / "billmann2026"

# ---------------------------------------------------------------------------
# Plant + yeast: canonical-PPI tuples → weighted edges with weight=1.0
# ---------------------------------------------------------------------------

for species, edges in [
    ("yeast", YEAST_CANONICAL_EDGES),
    ("arabidopsis", ARABIDOPSIS_CANONICAL_EDGES),
    ("rice", RICE_CANONICAL_EDGES),
]:
    rows = edges_from_canonical_ppi(edges, weight=1.0)
    out_path = OUT / f"{species}.tsv"
    write_edges_tsv(rows, out_path)
    print(f"  {species}: {len(rows)} edges → {out_path}")

# ---------------------------------------------------------------------------
# Human: Billmann qGI HC + DepMap ED, optionally + canonical PPI fallback
# ---------------------------------------------------------------------------

print("\n--- human ---")
human_rows: list = []

billmann_qgi_tsv = BILLMANN_DIR / "billmann_hc_gi_edges.tsv"
if billmann_qgi_tsv.exists():
    qgi_rows = edges_from_billmann_qgi_tsv(billmann_qgi_tsv,
                                              use_stringent_only=False,
                                              qgi_scale=1.0)
    print(f"  billmann_qgi: {len(qgi_rows):,} (standard threshold)")
    human_rows.extend(qgi_rows)
else:
    print(f"  ⚠ {billmann_qgi_tsv.name} missing — skipping qGI source")

depmap_ed_tsv = BILLMANN_DIR / "billmann_depmap_ed_edges.tsv"
if depmap_ed_tsv.exists():
    ed_rows = edges_from_depmap_ed_tsv(depmap_ed_tsv, ed_scale=0.5)
    print(f"  depmap_ed: {len(ed_rows):,}")
    human_rows.extend(ed_rows)
else:
    print(f"  ⚠ {depmap_ed_tsv.name} missing — skipping ED source")

# Note: there is no HUMAN_CANONICAL_EDGES in canonical_ppi.py yet.
# When added (paper #2 revision target), include it here:
# from graphgwas.canonical_ppi import HUMAN_CANONICAL_EDGES
# human_rows.extend(edges_from_canonical_ppi(HUMAN_CANONICAL_EDGES, weight=1.0))

write_edges_tsv(human_rows, OUT / "human.tsv")
print(f"  human total: {len(human_rows):,} edge-source rows → {OUT / 'human.tsv'}")

# Quick load-back check
from graphgwas.weighted_substrate import load_weighted_edges, edge_diagnostics
for species in ["yeast", "arabidopsis", "rice", "human"]:
    edges = load_weighted_edges(OUT / f"{species}.tsv")
    diag = edge_diagnostics(edges)
    print(f"\n  load-back {species}: {diag.get('n_edges', 0):,} unique gene pairs, "
          f"{diag.get('n_source_observations', 0):,} source rows; "
          f"sources={diag.get('source_counts')}")
